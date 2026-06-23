#!/usr/bin/env python3
"""One-off, idempotent importer for the 2026 schedule archive."""

from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, select

import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app import (  # noqa: E402
    AircrewArrival,
    Assignment,
    Employee,
    OccupancySnapshot,
    Section,
    SessionLocal,
    Week,
    _extract_aircrew_times_from_cell,
)


SECTION_LABELS = {
    "FRONT DESK": "Front Desk",
    "SHUTTLE DRIVERS": "Shuttle",
    "BREAKFAST BAR": "Breakfast Bar",
    "MAINTENANCE": "Maintenance",
    "MANAGER": "Manager",
}

DISPLAY_NAMES = {
    "KC": "KC",
    "TONY H": "Tony H",
}

IGNORED_ROW_MARKERS = {
    "OCC",
    "OCCUPANCY",
    "!OCC / ADR",
    "ADR",
}


def clean_name(raw: object) -> str:
    text = re.sub(r"\*+$", "", str(raw or "").strip()).strip()
    upper = text.upper()
    return DISPLAY_NAMES.get(upper, text.title())


def normalized_shift(raw: object, section: str) -> str:
    text = str(raw or "").strip()
    text = re.sub(r"\s+", " ", text)
    upper = text.upper()
    if not text or upper in {"-", "--"}:
        return "Set"
    if "REQ VAC" in upper:
        return "REQ VAC"
    if "REQ OFF" in upper or "NO AM" in upper:
        return "TIME OFF"
    if upper == "OFF":
        return "OFF"
    if upper in {"N/A", "NA"}:
        return "N/A"
    if section == "Maintenance" and upper in {"IN", "0800 - 1630", "0800-1630"}:
        return "8AM–4:30PM"

    compact = upper.replace(" ", "").replace("–", "-").replace("—", "-")
    compact = re.sub(r"\([^)]*\)$", "", compact)
    # Three June source cells say 10pm-6pm, but the surrounding shuttle row,
    # the published database values, and the other four days all say 10am-6pm.
    if section == "Shuttle" and compact == "10:00PM-6:00PM":
        return "10:00am - 6:00pm"
    mappings = {
        "5AM-12PM": "5AM–12PM",
        "6AM-12PM": "6AM–12PM",
        "7AM-12PM": "7AM–12PM",
        "6AM-2PM": "AM (6:00AM–2:00PM)",
        "6:15AM-2:15PM": "AM (6:15AM–2:15PM)",
        "2PM-10PM": "PM (2:00PM–10:00PM)",
        "2:15PM-10:15PM": "PM (2:15PM–10:15PM)",
        "10PM-6AM": "Audit (10:00PM–6:00AM)",
        "10:15PM-6:15AM": "Audit (10:15PM–6:15AM)",
        "3:30AM-11:30AM": "AM (3:30AM–11:30AM)",
        "10:30AM-6:30PM": "Midday (10:30AM–6:30PM)",
        "5:30PM-1:30AM": "PM (5:30PM–1:30AM)",
        "5:45PM-1:45AM": "Crew (5:45PM–1:45AM)",
        "8PM-12AM": "Crew (8:00PM–12:00AM)",
        "9PM-1AM": "Crew (9:00PM–1:00AM)",
    }
    if compact in mappings:
        return mappings[compact]

    # Keep unusual genuine shifts, but remove transfer/training notes.
    return re.sub(r"\s*\([^)]*\)\s*$", "", text).strip()


def workbook_dates(sheet) -> tuple[int, list[date]]:
    for row in range(1, min(sheet.max_row, 10) + 1):
        values = [sheet.cell(row, col).value for col in range(5, 12)]
        if all(isinstance(v, (date, datetime)) for v in values):
            dates = [v.date() if isinstance(v, datetime) else v for v in values]
            # The January files were copied from a prior-year template. Their
            # weekday headers and filenames unambiguously identify January 2026.
            dates = [d.replace(year=2026) if d.year == 2025 else d for d in dates]
            return row, dates
    raise ValueError(f"Could not find seven schedule dates in {sheet.title}")


def parse_workbook(path: Path) -> dict:
    sheet = load_workbook(path, data_only=True).active
    _, dates = workbook_dates(sheet)
    section = None
    employees: list[tuple[str, str, list[str]]] = []
    occupancy: list[float | None] | None = None
    aircrew: dict[str, list[list[str]]] = {}

    for row in range(1, sheet.max_row + 1):
        section_cell = str(sheet.cell(row, 1).value or "").strip().upper()
        if section_cell in SECTION_LABELS:
            section = SECTION_LABELS[section_cell]
        raw_name = str(sheet.cell(row, 4).value or "").strip()
        marker = re.sub(r"^\*+", "", raw_name).strip().upper()
        values = [sheet.cell(row, col).value for col in range(5, 12)]
        if marker in IGNORED_ROW_MARKERS:
            occupancy = []
            for value in values:
                if value in (None, ""):
                    occupancy.append(None)
                elif isinstance(value, (int, float)):
                    occupancy.append(float(value) * 100 if float(value) <= 1 else float(value))
                else:
                    match = re.search(r"\d+(?:\.\d+)?", str(value))
                    occupancy.append(float(match.group()) if match else None)
            continue
        if marker in {"AEROMEX", "AEROMEXICO", "SKYWEST"}:
            carrier = "Aeromexico" if marker.startswith("AEROMEX") else "Skywest"
            aircrew[carrier] = [_extract_aircrew_times_from_cell(value) for value in values]
            continue
        if section in {"Front Desk", "Shuttle", "Breakfast Bar", "Maintenance"}:
            name = clean_name(raw_name)
            if (
                name
                and not marker.startswith("*")
                and marker not in IGNORED_ROW_MARKERS
                and any(value not in (None, "") for value in values)
            ):
                employees.append((section, name, [normalized_shift(v, section) for v in values]))

    return {
        "path": path,
        "start": dates[0],
        "dates": dates,
        "employees": employees,
        "occupancy": occupancy,
        "aircrew": aircrew,
    }


MARCH_05 = {
    "Front Desk": {
        "Emilyn": ["Audit (10:00PM–6:00AM)"] * 4 + ["Set", "Set", "Audit (10:00PM–6:00AM)"],
        "Abdi": ["Set", "Set", "Audit (10:15PM–6:15AM)", "Audit (10:15PM–6:15AM)", "Audit (10:00PM–6:00AM)", "Audit (10:00PM–6:00AM)", "Audit (10:15PM–6:15AM)"],
        "Raphael": ["Audit (10:15PM–6:15AM)", "Audit (10:15PM–6:15AM)", "Set", "PM (2:15PM–10:15PM)", "Set", "Set", "Set"],
        "Oscar": ["Set", "PM (2:15PM–10:15PM)", "PM (2:15PM–10:15PM)", "Set", "Audit (10:15PM–6:15AM)", "Audit (10:15PM–6:15AM)", "Set"],
        "Cindy": ["AM (6:00AM–2:00PM)", "Set", "Set", "AM (6:00AM–2:00PM)", "AM (6:00AM–2:00PM)", "AM (6:00AM–2:00PM)", "AM (6:00AM–2:00PM)"],
        "KC": ["Set", "AM (6:00AM–2:00PM)", "AM (6:00AM–2:00PM)", "AM (6:15AM–2:15PM)", "Set", "Set", "Set"],
        "Christian": ["AM (6:15AM–2:15PM)", "AM (6:15AM–2:15PM)", "AM (6:15AM–2:15PM)", "Set", "Set", "Set", "AM (6:15AM–2:15PM)"],
        "Ryan": ["PM (2:15PM–10:15PM)", "Set", "Set", "Set", "AM (6:15AM–2:15PM)", "AM (6:15AM–2:15PM)", "PM (2:15PM–10:15PM)"],
        "Troy": ["Set", "Set", "TIME OFF", "TIME OFF", "PM (2:00PM–10:00PM)", "PM (2:00PM–10:00PM)", "OFF"],
        "Terry": ["PM (2:00PM–10:00PM)", "Set", "Set", "Set", "PM (2:15PM–10:15PM)", "PM (2:15PM–10:15PM)", "PM (2:00PM–10:00PM)"],
        "Tristan": ["OFF", "Set", "Set", "Set", "Set", "OFF", "Set"],
        "Sato": ["Set"] * 7,
        "Brian": ["Set", "PM (2:00PM–10:00PM)", "PM (2:00PM–10:00PM)", "PM (2:00PM–10:00PM)", "Set", "Set", "Set"],
        "Ian": ["Set"] * 7,
        "Sara": ["Set"] * 7,
        "Jordan": ["Set"] * 7,
    },
    "Shuttle": {
        "Leo": ["Midday (10:30AM–6:30PM)", "Set", "Set", "Set", "Midday (10:30AM–6:30PM)", "Midday (10:30AM–6:30PM)", "Midday (10:30AM–6:30PM)"],
        "Tony": ["PM (5:30PM–1:30AM)", "Midday (10:30AM–6:30PM)", "Midday (10:30AM–6:30PM)", "Midday (10:30AM–6:30PM)", "Set", "Set", "PM (5:30PM–1:30AM)"],
        "Kevin": ["OFF", "PM (5:30PM–1:30AM)", "PM (5:30PM–1:30AM)", "PM (5:30PM–1:30AM)", "PM (5:30PM–1:30AM)", "OFF", "Set"],
        "John": ["TIME OFF", "TIME OFF", "TIME OFF", "TIME OFF", "TIME OFF", "Set", "Set"],
        "Michael": ["Set", "AM (3:30AM–11:30AM)", "Set", "Set", "Set", "Set", "Set"],
        "Tony H": ["AM (3:30AM–11:30AM)", "Set", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)", "Set", "Set"],
        "Toshiyuki": ["Set"] * 7,
        "Ed": ["Set"] * 7,
        "Robenson": ["Set", "Set", "Set", "Set", "Set", "PM (5:30PM–1:30AM)", "Set"],
        "Ming": ["Set", "Set", "Set", "Set", "Set", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)"],
    },
    "Breakfast Bar": {
        "Rose": ["5AM–12PM", "5AM–12PM", "5AM–12PM", "OFF", "OFF", "5AM–12PM", "5AM–12PM"],
        "Yoko": ["Set", "Set", "Set", "5AM–12PM", "5AM–12PM", "REQ VAC", "Set"],
        "Eurielle": ["Set", "6AM–12PM", "N/A", "N/A", "6AM–12PM", "N/A", "6AM–12PM"],
        "Anna": ["7AM–12PM", "7AM–12PM", "6AM–12PM", "6AM–12PM", "Set", "7AM–12PM", "Set"],
        "Merve": ["6AM–12PM", "N/A", "Set", "OFF", "7AM–12PM", "6AM–12PM", "7AM–12PM"],
        "Ayako": ["Set", "Set", "7AM–12PM", "7AM–12PM", "Set", "Set", "Set"],
    },
    "Maintenance": {
        "Ricardo": ["8AM–4:30PM", "Set", "Set", "8AM–4:30PM", "8AM–4:30PM", "8AM–4:30PM", "8AM–4:30PM"],
        "Oscar": ["Set"] * 7,
    },
}


PDF_TONY_H = {
    date(2026, 2, 19): ["Set", "Set", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)"],
    date(2026, 2, 26): ["AM (3:30AM–11:30AM)", "Set", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)", "Set", "Set"],
    date(2026, 3, 5): MARCH_05["Shuttle"]["Tony H"],
    date(2026, 3, 12): ["AM (3:30AM–11:30AM)", "Set", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)", "Set", "Set"],
    date(2026, 3, 19): ["AM (3:30AM–11:30AM)", "Set", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)", "Set", "Set"],
    date(2026, 3, 26): ["AM (3:30AM–11:30AM)", "Set", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)", "AM (3:30AM–11:30AM)", "Set", "Set"],
}

PDF_OCCUPANCY = {
    date(2026, 2, 19): [63.55, 81.31, 77.57, 58.88, 30.84, 35.51, 39.25],
    date(2026, 2, 26): [42.06, 64.49, 81.31, 61.68, 36.45, 53.27, 57.94],
    date(2026, 3, 5): [56.07, 71.03, 74.77, 57.94, 49.53, 41.12, 30.84],
    date(2026, 3, 12): [37.38, 62.62, 60.75, 42.99, 45.79, 39.25, 39.25],
    date(2026, 3, 19): [39.25, 57.94, 54.21, 50.47, 40.19, 42.99, 44.86],
    date(2026, 3, 26): [34.58, 55.14, 60.75, 33.64, 17.76, 23.36, None],
    date(2026, 4, 2): [59.81, 74.77, 82.24, 64.49, 57.01, 44.86, 55.14],
    date(2026, 4, 9): [54.21, 73.83, 71.96, 53.27, 37.38, 39.25, 37.38],
    date(2026, 4, 16): [38.32, 71.96, 69.16, 55.14, 42.06, 46.73, 44.86],
    date(2026, 4, 23): [44.86, 71.03, 67.29, 50.47, 35.51, 28.97, 41.12],
}


def get_or_create_employee(session, name: str, section_name: str) -> Employee:
    section = session.scalar(select(Section).where(Section.name == section_name))
    employee = session.scalar(
        select(Employee).where(Employee.name == name, Employee.section_id == section.id)
    )
    if employee is None:
        first, _, last = name.partition(" ")
        employee = Employee(
            name=name,
            first_name=first,
            last_name=last or None,
            section_id=section.id,
            is_active=False,
        )
        session.add(employee)
        session.flush()
    return employee


def replace_employee_week(session, week: Week, employee: Employee, dates: list[date], values: list[str]) -> None:
    session.execute(
        delete(Assignment).where(
            Assignment.week_id == week.id,
            Assignment.employee_id == employee.id,
        )
    )
    session.add_all(
        Assignment(week_id=week.id, employee_id=employee.id, date=d, value=v)
        for d, v in zip(dates, values)
    )


def import_archive(archive_dir: Path, dry_run: bool) -> None:
    parsed = [parse_workbook(path) for path in sorted(archive_dir.glob("*.xlsx"))]
    former = {
        ("Front Desk", "Sara"),
        ("Front Desk", "Jordan"),
        ("Shuttle", "Victor"),
        ("Shuttle", "Tony H"),
        ("Maintenance", "Wilfredo"),
    }
    with SessionLocal() as session:
        imported_weeks = []
        for item in parsed:
            week = session.scalar(select(Week).where(Week.start_date == item["start"]))
            if week is None:
                week = Week(start_date=item["start"])
                session.add(week)
                session.flush()
            dates = item["dates"]
            source_ids = []
            for section_name, name, values in item["employees"]:
                employee = get_or_create_employee(session, name, section_name)
                source_ids.append(employee.id)
                replace_employee_week(session, week, employee, dates, values)
            if source_ids:
                session.execute(
                    delete(Assignment).where(
                        Assignment.week_id == week.id,
                        Assignment.employee_id.not_in(source_ids),
                    )
                )
            if item["occupancy"] is not None:
                session.execute(delete(OccupancySnapshot).where(OccupancySnapshot.week_id == week.id))
                session.add_all(
                    OccupancySnapshot(week_id=week.id, date=d, percentage=p, uploaded_at=datetime.utcnow())
                    for d, p in zip(dates, item["occupancy"])
                    if p is not None
                )
            for carrier, day_times in item["aircrew"].items():
                session.execute(
                    delete(AircrewArrival).where(
                        AircrewArrival.week_id == week.id,
                        AircrewArrival.carrier == carrier,
                    )
                )
                session.add_all(
                    AircrewArrival(week_id=week.id, carrier=carrier, date=d, times=json.dumps(times))
                    for d, times in zip(dates, day_times)
                    if times
                )
            imported_weeks.append(item["start"].isoformat())

        # Replace the corrupted March 5 week directly from the scanned source.
        march_start = date(2026, 3, 5)
        march_week = session.scalar(select(Week).where(Week.start_date == march_start))
        if march_week is None:
            march_week = Week(start_date=march_start)
            session.add(march_week)
            session.flush()
        session.execute(delete(Assignment).where(Assignment.week_id == march_week.id))
        march_dates = [march_start + timedelta(days=i) for i in range(7)]
        for section_name, people in MARCH_05.items():
            for name, values in people.items():
                employee = get_or_create_employee(session, name, section_name)
                replace_employee_week(session, march_week, employee, march_dates, values)
        # The PDF pages are scanned images, so their occupancy rows are
        # transcribed explicitly instead of relying on lossy OCR.
        for start, values in PDF_OCCUPANCY.items():
            week = session.scalar(select(Week).where(Week.start_date == start))
            session.execute(delete(OccupancySnapshot).where(OccupancySnapshot.week_id == week.id))
            session.add_all(
                OccupancySnapshot(
                    week_id=week.id,
                    date=start + timedelta(days=offset),
                    percentage=value,
                    uploaded_at=datetime.utcnow(),
                )
                for offset, value in enumerate(values)
                if value is not None
            )

        # PDF-only historical rows for employees absent from the current roster.
        for start, values in PDF_TONY_H.items():
            week = session.scalar(select(Week).where(Week.start_date == start))
            employee = get_or_create_employee(session, "Tony H", "Shuttle")
            replace_employee_week(
                session,
                week,
                employee,
                [start + timedelta(days=i) for i in range(7)],
                values,
            )
        for start in PDF_TONY_H:
            week = session.scalar(select(Week).where(Week.start_date == start))
            for name in ("Sara", "Jordan"):
                employee = get_or_create_employee(session, name, "Front Desk")
                replace_employee_week(
                    session,
                    week,
                    employee,
                    [start + timedelta(days=i) for i in range(7)],
                    ["Set"] * 7,
                )

        for section_name, name in former:
            employee = get_or_create_employee(session, name, section_name)
            employee.is_active = False

        print(f"Parsed {len(parsed)} workbooks: {', '.join(imported_weeks)}")
        print("Former employees archived: Jordan, Sara, Tony H, Victor, Wilfredo")
        if dry_run:
            session.rollback()
            print("Dry run complete; no database changes were committed.")
        else:
            session.commit()
            print("Import committed.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("archive_dir", type=Path)
    parser.add_argument("--commit", action="store_true")
    args = parser.parse_args()
    import_archive(args.archive_dir, dry_run=not args.commit)
